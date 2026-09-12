"""
Tests for Phase 8: Multimodal Intelligence, Web Search, Documents (PDF/Word), and Media Transcription.
"""

import os
import tempfile
import pytest
from pathlib import Path
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from app.tools.sre_tools import (
    search_web_or_docs,
    inspect_document,
    export_incident_report,
    transcribe_media_recording
)
from app.services.orchestrator import agent_orchestrator
from app.core.session import OperatorSession, current_session


@pytest.fixture(autouse=True)
def bind_test_session():
    """Ensure each test runs with an authenticated SRE_COMMANDER session."""
    session = OperatorSession(authenticated=True, role="SRE_COMMANDER")
    token = current_session.set(session)
    try:
        yield session
    finally:
        current_session.reset(token)



def test_web_search_knowledge_and_fallback():
    """Verify search_web_or_docs retrieves structured technical documentation and URLs."""
    res_pg = search_web_or_docs("PostgreSQL connection pool max connections")
    assert res_pg["status"] == "success"
    assert res_pg["result_count"] > 0
    assert any("postgresql" in r["url"].lower() or "postgres" in r["title"].lower() for r in res_pg["results"])

    res_aws = search_web_or_docs("AWS status us-east-1")
    assert res_aws["status"] == "success"
    assert res_aws["result_count"] > 0
    assert any("aws" in r["url"].lower() or "amazon" in r["url"].lower() for r in res_aws["results"])

    res_generic = search_web_or_docs("Some obscure internal custom metric XYZ999")
    assert res_generic["status"] == "success"
    assert len(res_generic["results"]) > 0


def test_pdf_document_inspection_and_citations():
    """Verify inspect_document parses PDF pages and returns grounded page citations."""
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf_path = f.name

    try:
        doc = SimpleDocTemplate(pdf_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = [
            Paragraph("Emergency PostgreSQL Failover Runbook", styles['Heading1']),
            Paragraph("Step 1: Check replication lag on read-replica-02.", styles['Normal']),
            Paragraph("Step 2: Issue pg_ctl promote on the standby node.", styles['Normal'])
        ]
        doc.build(story)

        # Query specific keyword
        res = inspect_document(pdf_path, query="promote")
        assert res["status"] == "success"
        assert res["file_type"] == "pdf"
        assert res["matched_count"] > 0
        excerpt = res["excerpts"][0]
        assert "Page 1" in excerpt["citation"]
        assert "promote" in excerpt["snippet"].lower()

        # Non-matching query
        res_none = inspect_document(pdf_path, query="nonexistent_xyz")
        assert res_none["status"] == "success"
        assert res_none["matched_count"] == 0
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)


def test_docx_document_inspection_and_citations():
    """Verify inspect_document parses Word (.docx) documents with section citations."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        docx_path = f.name

    try:
        doc = docx.Document()
        doc.add_heading("Database Architecture Specification", level=1)
        doc.add_paragraph("The order-db cluster runs PostgreSQL 15 with Patroni auto-failover.")
        doc.add_heading("Connection Pooling Guidelines", level=2)
        doc.add_paragraph("PgBouncer is configured with a default pool size of 50 connections.")
        doc.save(docx_path)

        res = inspect_document(docx_path, query="PgBouncer")
        assert res["status"] == "success"
        assert res["file_type"] == "docx"
        assert res["matched_count"] > 0
        excerpt = res["excerpts"][0]
        assert "§Connection Pooling Guidelines" in excerpt["citation"]
        assert "PgBouncer" in excerpt["snippet"]
    finally:
        if os.path.exists(docx_path):
            os.remove(docx_path)


def test_export_incident_report_pdf_and_docx():
    """Verify export_incident_report produces valid binary PDF and Word (.docx) files."""
    # Test PDF Export
    pdf_res = export_incident_report(format="pdf", filename="test_incident_report.pdf")
    assert pdf_res["status"] == "success"
    assert pdf_res["format"] == "pdf"
    assert os.path.exists(pdf_res["file_path"])
    with open(pdf_res["file_path"], "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-"
    if os.path.exists(pdf_res["file_path"]):
        os.remove(pdf_res["file_path"])

    # Test DOCX Export
    docx_res = export_incident_report(format="docx", filename="test_incident_report.docx")
    assert docx_res["status"] == "success"
    assert docx_res["format"] == "docx"
    assert os.path.exists(docx_res["file_path"])
    with open(docx_res["file_path"], "rb") as f:
        header = f.read(4)
        assert header == b"PK\x03\x04"  # Standard zip signature for .docx
    if os.path.exists(docx_res["file_path"]):
        os.remove(docx_res["file_path"])


def test_transcribe_media_recording_simulation():
    """Verify transcribe_media_recording parses audio and video files with speaker diarization and chapters."""
    with tempfile.NamedTemporaryFile(suffix=".mp4", delete=False) as f:
        f.write(b"dummy_mp4_bytes")
        video_path = f.name

    try:
        res = transcribe_media_recording(video_path, media_type="video")
        assert res["status"] == "success"
        assert res["media_type"] == "video"
        assert res["duration_seconds"] > 0
        assert len(res["utterances"]) >= 2
        assert len(res["chapters"]) >= 1
        assert "Speaker A" in res["utterances"][0]["speaker"]
    finally:
        if os.path.exists(video_path):
            os.remove(video_path)


@pytest.mark.asyncio
async def test_orchestrator_multimodal_intent_routing():
    """Verify natural voice prompts route directly to multimodal tools in orchestrator."""
    # Web search prompt
    spoken, tools, _ = await agent_orchestrator.process_user_turn("J.A.R.V.I.S., search web for PostgreSQL connection pool")
    assert len(tools) == 1
    assert tools[0]["tool_name"] == "search_web_or_docs"
    assert "postgresql" in tools[0]["arguments"]["query"].lower()

    # Document inspection prompt
    spoken_doc, tools_doc, _ = await agent_orchestrator.process_user_turn("J.A.R.V.I.S., read document docs/ARCHITECTURE.md")
    assert len(tools_doc) == 1
    assert tools_doc[0]["tool_name"] == "inspect_document"

    # Export document prompt
    spoken_exp, tools_exp, _ = await agent_orchestrator.process_user_turn("J.A.R.V.I.S., export report as pdf")
    assert len(tools_exp) == 1
    assert tools_exp[0]["tool_name"] == "export_incident_report"
    assert tools_exp[0]["arguments"]["format"] == "pdf"

    # Transcribe media recording prompt
    spoken_tr, tools_tr, _ = await agent_orchestrator.process_user_turn("J.A.R.V.I.S., transcribe recording data/incident_call.wav")
    assert len(tools_tr) == 1
    assert tools_tr[0]["tool_name"] == "transcribe_media_recording"
